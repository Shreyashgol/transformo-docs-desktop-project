from __future__ import annotations

import logging
from pathlib import Path

import docx

from app.extraction.base import BaseExtractor
from app.extraction.models import ExtractionResult

logger = logging.getLogger("transformo.extraction.docx")


class DocxExtractor(BaseExtractor):
    """Extracts text and metadata from DOCX files using python-docx."""

    def extract(self, path: Path) -> ExtractionResult:
        try:
            doc = docx.Document(str(path))
        except Exception as exc:
            raise ValueError(f"Failed to open DOCX document: {exc}") from exc

        parts: list[str] = []

        # Extract paragraph text
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                parts.append(text)

        # Extract table text
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    parts.append(" | ".join(row_cells))

        combined_text = "\n\n".join(parts).strip()

        # Extract core properties safely
        metadata = {"format": "DOCX"}
        try:
            core = doc.core_properties
            if core.title:
                metadata["title"] = core.title
            if core.author:
                metadata["author"] = core.author
            if core.subject:
                metadata["subject"] = core.subject
            if core.created:
                metadata["creation_date"] = core.created.isoformat()
            if core.modified:
                metadata["mod_date"] = core.modified.isoformat()
        except Exception as prop_err:
            logger.debug("Could not read DOCX core properties: %s", prop_err)

        # Estimate page count (standard rule of thumb: ~400 words or 2500 chars per page, min 1)
        estimated_pages = max(1, (len(combined_text) + 2000) // 2500) if combined_text else 1

        warnings: list[str] = []
        if not combined_text:
            warnings.append("Document is empty or contains no extractable text.")

        return ExtractionResult(
            text=combined_text,
            page_count=estimated_pages,
            metadata=metadata,
            warnings=warnings,
            is_scanned=False,
        )
