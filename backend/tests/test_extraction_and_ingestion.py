import io
from pathlib import Path
import docx
import fitz
import pytest

from app.extraction.docx import DocxExtractor
from app.extraction.pdf import PdfExtractor
from app.extraction.txt import TxtExtractor
from app.ingestion.service import (
    DuplicateDocumentError,
    EmptyFileError,
    FileSizeLimitError,
    InvalidFileContentError,
    UnsupportedFormatError,
)


def create_sample_pdf(text: str = "Hello Transformo Docs PDF") -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def create_scanned_like_pdf() -> bytes:
    # PDF with a page containing no text (simulating scanned document)
    doc = fitz.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def create_sample_docx(text: str = "Hello Transformo Docs DOCX") -> bytes:
    doc = docx.Document()
    doc.add_heading("Document Title", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_txt_extractor(tmp_path):
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Line 1\r\nLine 2\r\nLine 3", encoding="utf-8")

    extractor = TxtExtractor()
    result = extractor.extract(txt_file)
    assert "Line 1\nLine 2\nLine 3" == result.text
    assert result.metadata["format"] == "TXT"
    assert result.metadata["line_count"] == 3
    assert result.is_scanned is False


def test_pdf_extractor(tmp_path):
    pdf_bytes = create_sample_pdf("Extracted invoice contents #12345")
    pdf_file = tmp_path / "sample.pdf"
    pdf_file.write_bytes(pdf_bytes)

    extractor = PdfExtractor()
    result = extractor.extract(pdf_file)
    assert "Extracted invoice contents #12345" in result.text
    assert result.page_count == 1
    assert result.is_scanned is False
    assert len(result.warnings) == 0


def test_pdf_scanned_detection(tmp_path):
    pdf_bytes = create_scanned_like_pdf()
    pdf_file = tmp_path / "scanned.pdf"
    pdf_file.write_bytes(pdf_bytes)

    extractor = PdfExtractor()
    result = extractor.extract(pdf_file)
    assert result.is_scanned is True
    assert any("OCR is not yet available" in w for w in result.warnings)


def test_docx_extractor(tmp_path):
    docx_bytes = create_sample_docx("Employee Resume Content: Software Engineer")
    docx_file = tmp_path / "sample.docx"
    docx_file.write_bytes(docx_bytes)

    extractor = DocxExtractor()
    result = extractor.extract(docx_file)
    assert "Employee Resume Content: Software Engineer" in result.text
    assert result.metadata["format"] == "DOCX"


def test_api_import_txt(client, auth):
    content = b"Sample plain text file for testing."
    files = {"file": ("notes.txt", content, "text/plain")}
    res = client.post("/api/documents/import", headers=auth, files=files)
    assert res.status_code == 201
    body = res.json()
    assert body["metadata"]["original_filename"] == "notes.txt"
    assert body["metadata"]["extension"] == ".txt"
    assert "Sample plain text file" in body["extracted_text"]
    doc_id = body["metadata"]["id"]

    # Verify retrieval
    get_res = client.get(f"/api/documents/{doc_id}", headers=auth)
    assert get_res.status_code == 200
    assert get_res.json()["metadata"]["id"] == doc_id


def test_api_import_pdf(client, auth):
    pdf_bytes = create_sample_pdf("Testing PDF upload via API")
    files = {"file": ("report.pdf", pdf_bytes, "application/pdf")}
    res = client.post("/api/documents/import", headers=auth, files=files)
    assert res.status_code == 201
    body = res.json()
    assert body["metadata"]["original_filename"] == "report.pdf"
    assert body["metadata"]["page_count"] == 1
    assert "Testing PDF upload via API" in body["extracted_text"]


def test_api_import_docx(client, auth):
    docx_bytes = create_sample_docx("Testing DOCX upload via API")
    files = {"file": ("contract.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/documents/import", headers=auth, files=files)
    assert res.status_code == 201
    body = res.json()
    assert body["metadata"]["original_filename"] == "contract.docx"
    assert "Testing DOCX upload via API" in body["extracted_text"]


def test_api_import_duplicate_rejected(client, auth):
    content = b"Duplicate detection content unique 987654321."
    files = {"file": ("doc1.txt", content, "text/plain")}
    res1 = client.post("/api/documents/import", headers=auth, files=files)
    assert res1.status_code == 201

    # Second import with identical content
    files2 = {"file": ("doc2.txt", content, "text/plain")}
    res2 = client.post("/api/documents/import", headers=auth, files=files2)
    assert res2.status_code == 409
    body = res2.json()
    assert "detail" in body
    assert "already imported" in str(body["detail"])


def test_api_import_empty_rejected(client, auth):
    files = {"file": ("empty.txt", b"", "text/plain")}
    res = client.post("/api/documents/import", headers=auth, files=files)
    assert res.status_code == 400
    assert "empty file" in res.json()["detail"].lower()


def test_api_import_unsupported_format_rejected(client, auth):
    files = {"file": ("image.png", b"\x89PNG\r\n\x1a\n", "image/png")}
    res = client.post("/api/documents/import", headers=auth, files=files)
    assert res.status_code == 400
    assert "unsupported format" in res.json()["detail"].lower()


def test_api_import_malformed_pdf_rejected(client, auth):
    # Fake PDF extension with non-PDF content
    files = {"file": ("fake.pdf", b"NOT_A_REAL_PDF_HEADER", "application/pdf")}
    res = client.post("/api/documents/import", headers=auth, files=files)
    assert res.status_code == 422
    assert "valid pdf header" in res.json()["detail"].lower()


def test_api_list_documents(client, auth):
    res = client.get("/api/documents", headers=auth)
    assert res.status_code == 200
    body = res.json()
    assert "items" in body
    assert isinstance(body["items"], list)
    assert "total_count" in body


def test_api_get_nonexistent_document(client, auth):
    res = client.get("/api/documents/doc_nonexistent999", headers=auth)
    assert res.status_code == 404
